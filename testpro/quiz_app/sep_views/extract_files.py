import fitz  # PyMuPDF (PDF용)
import openpyxl  # Excel용
from pptx import Presentation  # PPT용
from docx import Document  # Word(DOCX)용
import os
# 파일 형식별 텍스트 추출 함수들
def extract_text_from_pptx(file_path):
    presentation = Presentation(file_path)
    text = [shape.text.strip() for slide in presentation.slides for shape in slide.shapes if
            hasattr(shape, "text") and shape.text.strip()]
    return "\n".join(text)


def extract_text_from_pdf(file_path):
    doc = fitz.open(file_path)
    text = [page.get_text("text") for page in doc]
    return "\n".join(text)


def extract_text_from_docx(file_path):
    doc = Document(file_path)
    text = [para.text.strip() for para in doc.paragraphs if para.text.strip()]
    return "\n".join(text)


def extract_text_from_xlsx(file_path):
    workbook = openpyxl.load_workbook(file_path)
    text = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            text.append(" | ".join([str(cell.value) for cell in row if cell.value]))
    return "\n".join(text)


def extract_text(file_path):
    ext = os.path.splitext(file_path)[1].lower()
    if ext == ".pptx":
        return extract_text_from_pptx(file_path)
    elif ext == ".pdf":
        return extract_text_from_pdf(file_path)
    elif ext == ".docx":
        return extract_text_from_docx(file_path)
    elif ext == ".xlsx":
        return extract_text_from_xlsx(file_path)
    else:
        raise ValueError("지원되지 않는 파일 형식입니다.")

