from django.shortcuts import render, redirect
from .forms import FileUploadForm
import openai
import fitz  # PyMuPDF (PDF용)
import openpyxl  # Excel용
from pptx import Presentation  # PPT용
from docx import Document  # Word(DOCX)용
from django.utils import timezone
from django.conf import settings
import re
import json
from .models import User_Quiz_Data
import os
from dotenv import load_dotenv  # dotenv 로드
from django.contrib.auth.decorators import login_required

# OpenAI API 클라이언트 초기화
client = openai.OpenAI(api_key=os.getenv("OPENAI_API_KEY"))

# 파일 형식별 텍스트 추출 함수들
def extract_text_from_pptx(file_path):
    presentation = Presentation(file_path)
    text = [shape.text.strip() for slide in presentation.slides for shape in slide.shapes if
            hasattr(shape, "text") and shape.text.strip()]
    return "\n".join(text)


def extract_text_from_pdf(file_path):
    doc = fitz.open(file_path)
    text = [page.get_text("text") for page in doc]
    return "\n".join(text)


def extract_text_from_docx(file_path):
    doc = Document(file_path)
    text = [para.text.strip() for para in doc.paragraphs if para.text.strip()]
    return "\n".join(text)


def extract_text_from_xlsx(file_path):
    workbook = openpyxl.load_workbook(file_path)
    text = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            text.append(" | ".join([str(cell.value) for cell in row if cell.value]))
    return "\n".join(text)


def extract_text(file_path):
    ext = os.path.splitext(file_path)[1].lower()
    if ext == ".pptx":
        return extract_text_from_pptx(file_path)
    elif ext == ".pdf":
        return extract_text_from_pdf(file_path)
    elif ext == ".docx":
        return extract_text_from_docx(file_path)
    elif ext == ".xlsx":
        return extract_text_from_xlsx(file_path)
    else:
        raise ValueError("지원되지 않는 파일 형식입니다.")


# GPT로 문제 생성 함수
def generate_quiz_with_gpt(text, question_type):
    system_prompt = """
    You have excellent knowledge in all fields. Create challenging and accurate problems based on the text provided by the user.
    """

    if question_type == "MCQ":
        prompt = f"""
            Based on the text below, please create 10 multiple-choice questions in Korean.

            === Given Text ===
            {text}

            📌 Formatting Requirements:
            - Each question must start with a number from "1." to "10.".
            - Each question must have exactly 4 answer choices, formatted as "a)", "b)", "c)", "d)".
            - Only one correct answer must be provided per question. Do not allow "all of the above" or multiple selections.
            - After the choices, write the answer line in the format: "정답: [answer letter]". The answer letter must be one of a, b, c, or d.
            - All answer choices must be distinct and logically consistent with the question.
            - Ensure that the answer provided exactly matches one of the choices.
            - For math problems, if necessary, use superscripts (e.g., 2³, x², 10⁵) and include any required functions or formulas.

            📌 Additional Requirements:
            - If the text contains mathematical formulas, theories, or important concepts, create applied problems that require understanding and problem-solving skills.
            - Use real-world examples to frame the questions wherever applicable.
            - For scientific and historical facts, design analytical or inferential questions rather than direct recall.
            - Ensure that problem difficulty is balanced, with a mix of direct application and higher-order thinking questions.

            📌 Pre-Submission Checklist:
            - Verify that the answer for each question logically matches the question and the provided choices.
            - Make sure there are no duplicate choices and no typos or logical errors in the questions or choices.
            - If a problem requires functions or mathematical expressions, ensure that such elements are correctly included.

            Example Format:
            1. [Applied Question based on given text]
            a) [Choice 1]
            b) [Choice 2]
            c) [Choice 3]
            d) [Choice 4]
            정답: [Correct answer letter]

            Please strictly adhere to this format when generating the questions.
        """
    elif question_type == "OX":
        prompt = f"""
            Based on the text below, please create 10 True/False (O/X) questions in Korean.

            === Given Text ===
            {text}

            📌 Formatting Requirements:
            - Each question must start with a number from "1." to "10.".
            - Write only the question text, followed by an answer line formatted as: "정답: [O or X]".
            - The answer must be either "O" or "X".
            - Verify that each answer logically corresponds to the question content.

            📌 Additional Requirements:
            - If the given text contains mathematical principles, physics laws, or technological concepts, modify them into applied questions rather than direct recall.
            - For scientific and historical information, test understanding through cause-and-effect-based true/false questions.

            Please strictly adhere to this format when generating the questions.
        """
    elif question_type == "Short":
        prompt = f"""
            Based on the text below, please create 10 short-answer questions in Korean.

            === Given Text ===
            {text}

            📌 Formatting Requirements:
            - Each question must start with a number from "1." to "10.".
            - Provide a single-line question followed by an answer line in the format: "정답: [Answer text]".
            - The answer must be a single word or a very short phrase (e.g., "타원", "산란") without any trailing punctuation or extra explanation.
            - Do not generate questions that ask for a process explanation (e.g., "과정을 설명하세요", "어떻게 ~하는지 서술하시오"). Instead, focus on fact-based questions answerable with one word or a short phrase.
            - Ensure that the question and answer logically match and are free of errors.

            📌 Additional Requirements:
            - If the given text contains formulas or theoretical principles, create problem-solving or applied questions that can be answered concisely.
            - Make sure questions require logical deduction or direct recall of facts, not extended explanations.
            - Ensure that answers are limited to one word or a very short phrase without any trailing punctuation.

            Please strictly adhere to this format when generating the questions.
        """
    else:
        raise ValueError("Unsupported question type.")

    response = client.chat.completions.create(
        model="gpt-4o-mini",
        messages=[
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": prompt}
        ],
        max_tokens=10000
    )

    return response.choices[0].message.content


# 문제 유형에 따라 문제와 답을 딕셔너리로 변환하는 함수
def parse_quiz_to_dict(quiz_text, question_type):
    quiz_dict = {}

    print(quiz_text)  # 디버깅용, 생성된 문제 원본을 확인하기 위해 출력

    if question_type == "MCQ":  # 객관식 문제
        question_pattern = r"(\d+)\.\s(.+?)(?=\n\d+\.|\Z)"
        choice_pattern = r"([a-d])\)\s(.+)"
        correct_answer_pattern = r"정답:\s*([a-d])"

        questions = re.findall(question_pattern, quiz_text, re.S)

        for idx, question_text in questions:
            choices = re.findall(choice_pattern, question_text)
            choices_dict = {choice[0]: choice[1].strip() for choice in choices if choice[0] in ['a', 'b', 'c', 'd']}
            correct_answer_match = re.search(correct_answer_pattern, question_text)
            correct_answer = correct_answer_match.group(1) if correct_answer_match else "오류"

            quiz_dict[idx] = {
                'question': question_text.strip().split("\n")[0],
                'choices': choices_dict,
                'correct_answer': correct_answer
            }

    elif question_type == "OX":  # OX 문제
        question_pattern = r"(\d+)\.\s(.+?)(?=\n\d+\.|\Z)"
        correct_answer_pattern = r"정답:\s*(O|X)"

        questions = re.findall(question_pattern, quiz_text, re.S)

        for idx, question_text in questions:
            correct_answer_match = re.search(correct_answer_pattern, question_text)
            correct_answer = correct_answer_match.group(1) if correct_answer_match else "오류"

            quiz_dict[idx] = {
                'question': question_text.strip().split("\n")[0],
                'choices': None,  # OX 문제는 선택지가 없음
                'correct_answer': correct_answer
            }

    elif question_type == "Short":  # 단답형 문제
        question_pattern = r"(\d+)\.\s(.+?)(?=\n\d+\.|\Z)"
        correct_answer_pattern = r"정답:\s*(.+)"

        questions = re.findall(question_pattern, quiz_text, re.S)

        for idx, question_text in questions:
            correct_answer_match = re.search(correct_answer_pattern, question_text)
            correct_answer = correct_answer_match.group(1) if correct_answer_match else "오류"

            quiz_dict[idx] = {
                'question': question_text.strip().split("\n")[0],
                'choices': None,  # 단답형 문제는 선택지가 없음
                'correct_answer': correct_answer
            }

    return quiz_dict

# db 데이터 저장
def save_quiz_to_db(quiz_dict, user, file_name, question_type):
    User_Quiz_Data.objects.create(
        created_at=timezone.now(),
        quiz_data=quiz_dict,
        user=user,
        file_name=file_name,  # 사용자가 지정한 파일 이름도 함께 저장
        question_type=question_type  # 문제 유형 저장
    )


def index(request):
    form = FileUploadForm()
    quiz_text = None  # GPT가 생성한 퀴즈 원본 텍스트
    if request.method == "POST":
        form = FileUploadForm(request.POST, request.FILES)
        if form.is_valid():
            file = form.cleaned_data["file"]
            question_type = form.cleaned_data["question_type"]  # 사용자가 선택한 문제 유형
            file_name = form.cleaned_data["file_name"] or "Untitled"  # 사용자가 입력한 파일 이름 없으면 "Untitled"로 설정
            upload_folder = "uploads"
            if not os.path.exists(upload_folder):
                os.makedirs(upload_folder)

            file_path = os.path.join(upload_folder, file.name)
            with open(file_path, "wb") as f:
                for chunk in file.chunks():
                    f.write(chunk)

            try:
                file_text = extract_text(file_path)
                if file_text:
                    # GPT로 문제 생성
                    quiz_text = generate_quiz_with_gpt(file_text, question_type)
                    quiz_dict = parse_quiz_to_dict(quiz_text, question_type)
                    save_quiz_to_db(quiz_dict, request.user, file_name, question_type)

                    # 문제 유형에 따라 다른 뷰로 리다이렉트
                    if question_type == "MCQ":
                        return redirect('multiple')
                    elif question_type == "Short":
                        return redirect('short')
                    else:
                        return redirect('OX')

                else:
                    return render(request, 'quiz_app/index.html', {
                        'error': "파일에서 텍스트를 추출할 수 없습니다.",
                        'form': form
                    })
            except Exception as e:
                return render(request, 'quiz_app/index.html', {
                    'error': f"오류 발생: {e}",
                    'form': form
                })

    return render(request, 'quiz_app/index.html', {'form': form})


def multiple_view(request):
    # 1. DB에서 최신 퀴즈 데이터 가져오기
    user = request.user
    quiz_data_obj = User_Quiz_Data.objects.filter(user=user).order_by('-created_at').first()

    if quiz_data_obj:
        quiz_dict = quiz_data_obj.quiz_data  # 예: {"1": {...}, "2": {...}, ...}
        question_type = quiz_data_obj.question_type
    else:
        quiz_dict = None
        question_type = None

    # 2. 현재 문제 번호 (1번부터 시작)
    current_question_number = int(request.GET.get('question_number', 1))

    # 3. 현재 문제 가져오기
    current_question = quiz_dict.get(str(current_question_number), None) if quiz_dict else None

    # 4. 템플릿에 전달할 컨텍스트
    context = {
        'current_question': current_question,
        'quiz_data': quiz_dict,
        'question_type': question_type,
        'current_question_number': current_question_number,
    }
    # Multiple Choice (객관식) → example.html
    return render(request, 'quiz_app/example.html', context)

def short_view(request):
    # 1. DB에서 최신 퀴즈 데이터 가져오기
    user = request.user
    quiz_data_obj = User_Quiz_Data.objects.filter(user=user).order_by('-created_at').first()

    if quiz_data_obj:
        quiz_dict = quiz_data_obj.quiz_data
        question_type = quiz_data_obj.question_type
    else:
        quiz_dict = {}
        question_type = None

    # 2. quiz_dict의 데이터를 프론트엔드에서 사용할 배열 형태로 변환
    questions = []
    # quiz_dict의 key가 문자열 숫자("1", "2", ...)이므로 정렬할 때 int로 변환합니다.
    for key in sorted(quiz_dict.keys(), key=int):
        q = quiz_dict[key]
        questions.append({
            'text': q.get('question', ''),
            'answer': q.get('correct_answer', '')
        })

    # 3. 템플릿에 전달할 컨텍스트 (questions 리스트를 포함)
    context = {
        'questions': questions,
        'quiz_data': quiz_dict,
        'question_type': question_type,
    }
    # Short Answer (단답형) → munje_shortanswer.html
    return render(request, 'quiz_app/munje_shortanswer.html', context)

def OX_view(request):
    # 1. DB에서 최신 퀴즈 데이터 가져오기
    user = request.user
    quiz_data_obj = User_Quiz_Data.objects.filter(user=user).order_by('-created_at').first()

    if quiz_data_obj:
        quiz_dict = quiz_data_obj.quiz_data  # 예: {"1": {"question": "...", "correct_answer": "O"}, ...}
        question_type = quiz_data_obj.question_type
    else:
        quiz_dict = {}
        question_type = None

    # 2. DB에 저장된 O/X 퀴즈 데이터를 프론트엔드에서 사용할 배열 형태로 변환
    # 각 문제는 { text, options, correctIndex } 형태가 되도록 변환합니다.
    # - options는 고정으로 ["참", "거짓"]입니다.
    # - correctIndex는 correct_answer가 "O"이면 0, "X"이면 1로 설정합니다.
    questions = []
    for key in sorted(quiz_dict.keys(), key=int):
        q = quiz_dict[key]
        correct_ans = q.get("correct_answer", "O")
        correctIndex = 0 if correct_ans.upper() == "O" else 1
        questions.append({
            "text": q.get("question", ""),
            "options": ["참", "거짓"],
            "correctIndex": correctIndex
        })

    # 3. 템플릿에 전달할 컨텍스트 구성
    context = {
        "questions": questions,
        "quiz_data": quiz_dict,
        "question_type": question_type,
    }
    # O/X (True/False) → munje_truefalse.html
    return render(request, "quiz_app/munje_truefalse.html", context)